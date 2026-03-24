#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = Path('/Users/chinnadurairamachandran/Downloads/Equipment Status.xlsx')
DEFAULT_OUTPUT = ROOT / 'output' / 'spreadsheet' / 'equipment_status_normalized.json'


def slugify(value: str, fallback: str = 'ASSET', max_length: int = 32) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9]+', '-', value.strip().upper()).strip('-')
    cleaned = re.sub(r'-{2,}', '-', cleaned)
    return cleaned[:max_length] or fallback


def parse_sheet_date(cell_value: Any) -> str | None:
    if isinstance(cell_value, datetime):
        return cell_value.date().isoformat()
    if not cell_value:
        return None
    text = str(cell_value)
    match = re.search(r'(\d{2})[./-](\d{2})[./-](\d{4})', text)
    if not match:
        return None
    day, month, year = match.groups()
    return f'{year}-{month}-{day}'


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_status(value: Any) -> str:
    text = (normalize_text(value) or '').upper()
    compact = re.sub(r'\s+', ' ', text)
    if not compact:
        return 'UNKNOWN'
    if 'UNDER MOH' in compact:
        return 'UNDER_MOH'
    if 'OVERHAUL' in compact:
        return 'OVERHAULING'
    if 'NOT WORK' in compact:
        return 'NOT_WORKING'
    if compact == 'WORKING' or 'WORKING' in compact:
        return 'WORKING'
    return 'UNKNOWN'


def infer_service_line(sheet_name: str) -> str:
    normalized = sheet_name.strip().lower()
    return 'WORKOVER' if 'workover' in normalized or 'well service' in normalized else 'SURFACE'


def infer_installation_type(sheet_name: str) -> str:
    normalized = sheet_name.strip().lower()
    return 'Workover Rig' if 'workover' in normalized or 'well service' in normalized else 'Surface'


def is_workover_sheet(sheet_name: str, worksheet) -> bool:
    normalized = sheet_name.strip().lower()
    if 'workover' in normalized or 'well service' in normalized:
        return True

    for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 6), values_only=True):
        joined = ' '.join(str(cell or '') for cell in row).lower()
        if 'installation/ rig' in joined or 'details of critical equipment of well services' in joined:
            return True

    return False


def infer_category(equipment_name: str) -> str:
    electrical_markers = ['motor', 'generator', 'starter', 'alternator', 'transformer', 'panel', 'mcc', 'switchgear', 'battery']
    lowered = equipment_name.lower()
    if any(marker in lowered for marker in electrical_markers):
        return 'ELECTRICAL'
    return 'MECHANICAL'


def infer_equipment_type(equipment_name: str) -> str:
    lowered = equipment_name.lower()
    if 'engine' in lowered:
        return 'Engine'
    if 'compressor' in lowered:
        return 'Compressor'
    if 'pump' in lowered:
        return 'Pump'
    if 'motor' in lowered:
        return 'Motor'
    if 'generator' in lowered:
        return 'Generator'
    if 'rotary table' in lowered:
        return 'Rotary Table'
    if 'starter' in lowered:
        return 'Starter'
    if 'fan' in lowered:
        return 'Fan'
    words = [word for word in re.split(r'[^A-Za-z0-9]+', equipment_name) if word]
    return ' '.join(words[:3]) or 'Equipment'


def build_tag_base(installation_name: str, equipment_name: str) -> str:
    installation_part = slugify(installation_name, fallback='INST', max_length=14)
    equipment_part = slugify(equipment_name, fallback='ASSET', max_length=28)
    return f'{installation_part}-{equipment_part}'


def parse_general_sheet(worksheet, sheet_name: str) -> tuple[list[dict[str, Any]], str | None]:
    sheet_date = parse_sheet_date(worksheet['B1'].value)
    rows: list[dict[str, Any]] = []
    current_installation: str | None = None

    for index, row in enumerate(worksheet.iter_rows(min_row=3, values_only=True), start=3):
        installation = normalize_text(row[0])
        equipment_name = normalize_text(row[1])
        make = normalize_text(row[2])
        model = normalize_text(row[3])
        serial_number = normalize_text(row[4])
        running_status = normalize_text(row[5])
        status_reason = normalize_text(row[6])

        if installation:
            current_installation = installation
        if not equipment_name or not current_installation:
            continue

        rows.append(
            {
                'sheetName': sheet_name,
                'sourceRow': index,
                'installationName': current_installation,
                'equipmentName': equipment_name,
                'make': make,
                'model': model,
                'serialNumber': serial_number,
                'assetCode': None,
                'operationalStatus': normalize_status(running_status),
                'statusReason': status_reason,
                'statusUpdatedAt': sheet_date,
                'serviceLine': infer_service_line(sheet_name),
                'installationType': infer_installation_type(sheet_name),
                'category': infer_category(equipment_name),
                'equipmentTypeName': infer_equipment_type(equipment_name),
                'rawStatus': running_status,
            }
        )

    return rows, sheet_date


def parse_workover_sheet(worksheet, sheet_name: str) -> tuple[list[dict[str, Any]], str | None]:
    sheet_date = parse_sheet_date(worksheet['A2'].value)
    rows: list[dict[str, Any]] = []
    current_installation: str | None = None
    header_row_index = 3

    for idx, row in enumerate(worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 8), values_only=True), start=1):
        joined = ' '.join(str(cell or '') for cell in row).lower()
        if 'installation/ rig' in joined:
            header_row_index = idx
            break

    for index, row in enumerate(worksheet.iter_rows(min_row=header_row_index + 2, values_only=True), start=header_row_index + 2):
        installation = normalize_text(row[1])
        equipment_name = normalize_text(row[2])
        asset_code = normalize_text(row[3])
        make = normalize_text(row[4])
        model = normalize_text(row[5])
        serial_number = normalize_text(row[6])
        running_status = normalize_text(row[7])
        status_reason = normalize_text(row[8])

        if installation:
            current_installation = installation
        if not equipment_name or not current_installation:
            continue

        rows.append(
            {
                'sheetName': sheet_name,
                'sourceRow': index,
                'installationName': current_installation,
                'equipmentName': equipment_name,
                'make': make,
                'model': model,
                'serialNumber': serial_number,
                'assetCode': asset_code,
                'operationalStatus': normalize_status(running_status),
                'statusReason': status_reason,
                'statusUpdatedAt': sheet_date,
                'serviceLine': infer_service_line(sheet_name),
                'installationType': infer_installation_type(sheet_name),
                'category': infer_category(equipment_name),
                'equipmentTypeName': infer_equipment_type(equipment_name),
                'rawStatus': running_status,
            }
        )

    return rows, sheet_date


def main() -> int:
    workbook_path = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    output_path = Path(sys.argv[2]).expanduser() if len(sys.argv) > 2 else DEFAULT_OUTPUT

    if not workbook_path.exists():
        raise FileNotFoundError(f'Workbook not found: {workbook_path}')

    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    parsed_rows: list[dict[str, Any]] = []
    sheet_summaries: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        if is_workover_sheet(worksheet.title, worksheet):
            rows, sheet_date = parse_workover_sheet(worksheet, worksheet.title)
        else:
            rows, sheet_date = parse_general_sheet(worksheet, worksheet.title)
        parsed_rows.extend(rows)
        sheet_summaries.append(
            {
                'sheetName': worksheet.title,
                'rowCount': len(rows),
                'sheetDate': sheet_date,
            }
        )

    tag_counter: Counter[str] = Counter()
    installations: dict[str, dict[str, Any]] = {}
    normalized_equipment: list[dict[str, Any]] = []

    for row in parsed_rows:
        base_tag = build_tag_base(row['installationName'], row['equipmentName'])
        tag_counter[base_tag] += 1
        equipment_tag = base_tag if tag_counter[base_tag] == 1 else f'{base_tag}-{tag_counter[base_tag]}'

        installations[row['installationName']] = {
            'installationId': row['installationName'],
            'location': row['sheetName'],
            'type': row['installationType'],
            'serviceLine': row['serviceLine'],
            'sourceSheet': row['sheetName'],
        }

        normalized_equipment.append(
            {
                'equipmentTag': equipment_tag,
                'installationId': row['installationName'],
                'description': row['equipmentName'],
                'make': row['make'],
                'model': row['model'],
                'serialNumber': row['serialNumber'],
                'assetCode': row['assetCode'],
                'category': row['category'],
                'equipmentTypeName': row['equipmentTypeName'],
                'serviceLine': row['serviceLine'],
                'operationalStatus': row['operationalStatus'],
                'statusReason': row['statusReason'],
                'statusUpdatedAt': row['statusUpdatedAt'],
                'sheetName': row['sheetName'],
                'sourceRow': row['sourceRow'],
                'rawStatus': row['rawStatus'],
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        'generatedAt': datetime.now(UTC).isoformat(),
        'sourceWorkbook': str(workbook_path),
        'sheetSummaries': sheet_summaries,
        'installations': list(installations.values()),
        'equipment': normalized_equipment,
    }
    output_path.write_text(json.dumps(payload, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(output_path), 'installations': len(installations), 'equipment': len(normalized_equipment)}, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
