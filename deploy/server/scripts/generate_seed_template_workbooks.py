#!/usr/bin/env python3
"""Generate seed-data workbook templates grouped by maintenance app module."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT_DIR = Path(__file__).resolve().parents[3]
SCHEMA_PATH = Path(__file__).resolve().parents[1] / "prisma" / "schema.prisma"
OUTPUT_DIR = ROOT_DIR / "output" / "spreadsheet" / "maintenance_seed_templates"

BUILTIN_TYPES = {"String", "Int", "Float", "Boolean", "DateTime", "Json", "Bytes", "Decimal", "BigInt"}

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
TITLE_FILL = PatternFill("solid", fgColor="0F172A")
REQUIRED_FILL = PatternFill("solid", fgColor="DBEAFE")
OPTIONAL_FILL = PatternFill("solid", fgColor="F3F4F6")
AUTO_FILL = PatternFill("solid", fgColor="FFEDD5")
README_HEADER_FILL = PatternFill("solid", fgColor="0F766E")


@dataclass
class FieldInfo:
    name: str
    prisma_type: str
    status: str
    notes: List[str] = field(default_factory=list)


@dataclass
class ModelInfo:
    name: str
    fields: List[FieldInfo] = field(default_factory=list)


WORKBOOKS: List[Tuple[str, str, List[str]]] = [
    (
        "01_admin_org_templates.xlsx",
        "Administration, users, organization, and budget masters",
        ["Company", "Department", "User", "Budget", "LaborRate"],
    ),
    (
        "02_procurement_templates.xlsx",
        "Procurement case and commentary seeding",
        ["Case", "CaseComment"],
    ),
    (
        "03_equipment_registry_templates.xlsx",
        "Equipment, instruments, calibration masters, and registry data",
        [
            "Installation",
            "InstrumentType",
            "MeterType",
            "EquipmentType",
            "Product",
            "CalibrationPerformer",
            "CalibrationStandard",
            "InstrumentMaster",
            "RunningEquipmentMaster",
            "CustodyTransferMeter",
            "InternalFlowMeter",
            "PMSchedule",
            "CalibrationLog",
            "ReplacementHistory",
            "PMSInitialState",
        ],
    ),
    (
        "04_asset_hierarchy_templates.xlsx",
        "ISO 14224 asset hierarchy, strategies, and location mapping",
        [
            "Site",
            "Area",
            "System",
            "FunctionalLocation",
            "Asset",
            "AssetInstallation",
            "MaintenanceStrategy",
            "StrategyTask",
            "MaintenancePlanAssignment",
            "OperationalReading",
            "SkidConfiguration",
            "SkidDriveTrain",
            "EngineRegistry",
            "FLAssetAssignment",
            "AssignmentAuditLog",
            "StaticEquipment",
        ],
    ),
    (
        "05_work_execution_templates.xlsx",
        "Maintenance requests, work orders, teams, and execution data",
        [
            "MaintenanceRequest",
            "WorkOrder",
            "WorkOrderChecklist",
            "WorkOrderTeam",
            "WOAttachment",
            "FailureMode",
            "FailureMechanism",
            "CauseCode",
            "ActionCode",
            "Notification",
        ],
    ),
    (
        "06_calibration_advanced_templates.xlsx",
        "Advanced calibration events, traceability, and reliability models",
        [
            "CalibrationEvent",
            "CalibrationPoint",
            "ISO14224Level",
            "RCMAnalysis",
            "CalibrationTraceability",
            "UnifiedCalibrationHistory",
        ],
    ),
    (
        "07_training_energy_operations_templates.xlsx",
        "Operations logs, manpower, training, energy, workshop, and MOH",
        [
            "Manpower",
            "MaintenanceLog",
            "MaintenanceLogTeam",
            "EquipmentLog",
            "TrainingRecord",
            "TrainingAttendee",
            "DailyEnergyLog",
            "MonthlyElectricityBill",
            "WorkshopJob",
            "MOHRecord",
        ],
    ),
    (
        "08_collaboration_templates.xlsx",
        "Collaboration, feedback, reactions, and presentation content",
        [
            "Discussion",
            "DiscussionReply",
            "DiscussionReaction",
            "ReplyReaction",
            "Feedback",
            "FeedbackReaction",
            "FeedbackComment",
            "Presentation",
        ],
    ),
    (
        "09_contracts_templates.xlsx",
        "Contracts, contractor access scopes, milestones, and reports",
        [
            "Contract",
            "ContractMilestone",
            "ContractDocument",
            "ContractUserAccess",
            "ContractInstallationScope",
            "ContractInstrumentTypeScope",
            "ContractEquipmentScope",
            "ContractFunctionalLocationScope",
            "WorkOrderExecutionLog",
            "ContractorReport",
        ],
    ),
    (
        "10_manuals_audit_templates.xlsx",
        "Manual repository, audit observations, process logs, and gas compression",
        [
            "ManualFolder",
            "RepositoryDocument",
            "AuditObservation",
            "AuditAction",
            "ActivityAuditLog",
            "GasCompressionLog",
        ],
    ),
    (
        "11_inspection_templates.xlsx",
        "Inspection rounds and execution records",
        ["InspectionRound", "InspectionExecution"],
    ),
]


def scalar_type(type_token: str) -> Tuple[str, bool]:
    is_array = type_token.endswith("[]")
    base = type_token.replace("?", "")
    if is_array:
        base = base[:-2]
    return base, is_array


def is_scalar(type_token: str) -> bool:
    base, _ = scalar_type(type_token)
    return base in BUILTIN_TYPES


def parse_schema(schema_path: Path) -> Dict[str, ModelInfo]:
    models: Dict[str, ModelInfo] = {}
    relation_hints: Dict[str, Dict[str, str]] = {}
    current_model: ModelInfo | None = None

    for raw_line in schema_path.read_text(encoding="utf-8").splitlines():
        model_match = re.match(r"^model\s+(\w+)\s+\{$", raw_line.strip())
        if model_match:
            current_model = ModelInfo(name=model_match.group(1))
            models[current_model.name] = current_model
            relation_hints[current_model.name] = {}
            continue

        if current_model is None:
            continue

        stripped = raw_line.strip()
        if stripped == "}":
            current_model = None
            continue

        if not stripped or stripped.startswith("//"):
            continue

        content, _, inline_comment = raw_line.partition("//")
        field_line = content.strip()
        if not field_line or field_line.startswith("@@"):
            continue

        field_match = re.match(r"^(\w+)\s+([^\s]+)\s*(.*)$", field_line)
        if not field_match:
            continue

        field_name, type_token, attributes = field_match.groups()
        comment = inline_comment.strip()
        base_type, is_array = scalar_type(type_token)

        if "@relation(" in attributes and not is_scalar(type_token):
            relation_match = re.search(
                r"@relation\((?:\"[^\"]+\"\s*,\s*)?fields:\s*\[([^\]]+)\],\s*references:\s*\[([^\]]+)\]",
                attributes,
            )
            if relation_match:
                local_fields = [part.strip() for part in relation_match.group(1).split(",")]
                referenced_fields = [part.strip() for part in relation_match.group(2).split(",")]
                for index, local_field in enumerate(local_fields):
                    ref_field = referenced_fields[min(index, len(referenced_fields) - 1)]
                    relation_hints[current_model.name][local_field] = f"Foreign key -> {base_type}.{ref_field}"
            continue

        if not is_scalar(type_token):
            continue

        auto_field = "@default(" in attributes or "@updatedAt" in attributes
        optional_field = "?" in type_token
        status = "AUTO" if auto_field else ("OPTIONAL" if optional_field else "REQUIRED")

        notes: List[str] = []
        if "@id" in attributes:
            notes.append("Primary key")
        if "@unique" in attributes:
            notes.append("Unique value")
        if is_array:
            notes.append("Enter multiple values using | as separator")
        if comment:
            notes.append(comment)

        current_model.fields.append(
            FieldInfo(
                name=field_name,
                prisma_type=type_token,
                status=status,
                notes=notes,
            )
        )

    for model_name, hints in relation_hints.items():
        for field in models[model_name].fields:
            hint = hints.get(field.name)
            if hint:
                field.notes.append(hint)

    return models


def set_column_widths(sheet) -> None:
    for column in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column[:8]:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 4, 14), 36)


def add_model_sheet(workbook: Workbook, model: ModelInfo) -> None:
    sheet = workbook.create_sheet(title=model.name[:31])
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"

    sheet["A1"] = model.name
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = TITLE_FILL
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(model.fields), 1))

    sheet["A2"] = "Populate data starting from row 7. Leave AUTO columns blank unless you need explicit seed values."
    sheet["A2"].font = Font(italic=True, color="475569")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(model.fields), 1))

    for column_index, field in enumerate(model.fields, start=1):
        header_cell = sheet.cell(row=4, column=column_index, value=field.name)
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.fill = HEADER_FILL
        header_cell.alignment = Alignment(horizontal="center", vertical="center")

        status_cell = sheet.cell(row=5, column=column_index, value=field.status)
        status_cell.alignment = Alignment(horizontal="center")
        if field.status == "REQUIRED":
            status_cell.fill = REQUIRED_FILL
        elif field.status == "OPTIONAL":
            status_cell.fill = OPTIONAL_FILL
        else:
            status_cell.fill = AUTO_FILL

        type_cell = sheet.cell(row=6, column=column_index, value=field.prisma_type)
        type_cell.font = Font(color="475569", italic=True)
        type_cell.alignment = Alignment(horizontal="center")

        if field.notes:
            header_cell.comment = Comment("\n".join(field.notes), "Codex")

    sheet.auto_filter.ref = f"A4:{get_column_letter(max(len(model.fields), 1))}4"
    sheet.row_dimensions[1].height = 24
    set_column_widths(sheet)


def add_readme_sheet(workbook: Workbook, filename: str, description: str, models: List[str], catalog: Dict[str, ModelInfo]) -> None:
    sheet = workbook.active
    sheet.title = "README"
    sheet.sheet_view.showGridLines = False

    sheet["A1"] = filename
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = TITLE_FILL
    sheet.merge_cells("A1:E1")

    sheet["A2"] = description
    sheet["A2"].font = Font(italic=True, color="475569")
    sheet.merge_cells("A2:E2")

    guidance = [
        "Rows 4-6 in each model sheet describe the columns.",
        "REQUIRED means seed data should normally be supplied.",
        "OPTIONAL means the column can be left blank.",
        "AUTO means Prisma/database normally generates or maintains the value.",
        "Header comments include relation and uniqueness hints when available.",
    ]
    for index, item in enumerate(guidance, start=4):
        sheet.cell(row=index, column=1, value=f"- {item}")

    start_row = 11
    headers = ["Order", "Sheet", "Required Columns", "Optional Columns", "Auto Columns"]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = README_HEADER_FILL

    for offset, model_name in enumerate(models, start=1):
        model = catalog[model_name]
        row = start_row + offset
        required = [field.name for field in model.fields if field.status == "REQUIRED"]
        optional = [field.name for field in model.fields if field.status == "OPTIONAL"]
        auto = [field.name for field in model.fields if field.status == "AUTO"]

        sheet.cell(row=row, column=1, value=offset)
        sheet.cell(row=row, column=2, value=model.name)
        sheet.cell(row=row, column=3, value=", ".join(required) or "-")
        sheet.cell(row=row, column=4, value=", ".join(optional) or "-")
        sheet.cell(row=row, column=5, value=", ".join(auto) or "-")

    for column in range(1, 6):
        sheet.column_dimensions[get_column_letter(column)].width = [8, 28, 44, 44, 44][column - 1]


def write_catalog(readme_path: Path, created_files: List[Tuple[str, str, List[str]]]) -> None:
    lines = [
        "# Maintenance Seed Template Catalog",
        "",
        "These workbooks were generated from the live Prisma schema in `deploy/server/prisma/schema.prisma`.",
        "",
    ]
    for filename, description, models in created_files:
        lines.append(f"## {filename}")
        lines.append("")
        lines.append(description)
        lines.append("")
        lines.append(f"Sheets: {', '.join(models)}")
        lines.append("")

    readme_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    models = parse_schema(SCHEMA_PATH)

    created_files: List[Tuple[str, str, List[str]]] = []
    assigned_models = set()

    for filename, description, workbook_models in WORKBOOKS:
        missing = [model_name for model_name in workbook_models if model_name not in models]
        if missing:
            raise ValueError(f"Missing models in schema for workbook {filename}: {', '.join(missing)}")

        workbook = Workbook()
        add_readme_sheet(workbook, filename, description, workbook_models, models)

        for model_name in workbook_models:
            assigned_models.add(model_name)
            add_model_sheet(workbook, models[model_name])

        workbook_path = OUTPUT_DIR / filename
        workbook.save(workbook_path)
        created_files.append((filename, description, workbook_models))
        print(f"Generated {workbook_path}")

    unassigned = sorted(set(models) - assigned_models)
    if unassigned:
        print("Unassigned models:", ", ".join(unassigned))

    write_catalog(OUTPUT_DIR / "README.md", created_files)


if __name__ == "__main__":
    main()
