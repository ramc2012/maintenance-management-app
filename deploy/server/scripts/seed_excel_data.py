#!/usr/bin/env python3
"""
Parse Equipment Status.xlsx and Instruments list Excel files,
then generate a TypeScript seed script for Prisma database seeding.

Usage:
    python3 seed_excel_data.py

Output:
    seed_excel_data.ts  (in same directory)
"""

import os
import re
import json
import uuid
from collections import OrderedDict
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

EQUIPMENT_FILE = "/Users/chinnadurairamachandran/Downloads/Equipment Status.xlsx"
INSTRUMENTS_FILE = "/Users/chinnadurairamachandran/Downloads/CMS CONTRACT/All instruments list isg 2025 (1).xlsx"
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "seed_excel_data.ts")

# ---------------------------------------------------------------------------
# Installation name normalization map
# ---------------------------------------------------------------------------
INSTALLATION_MAP = {
    # Area-1&2
    "GGS - 2": ("ANK-GGS-2", "GGS - 2, Ankleshwar", "Surface"),
    "GGS - 3": ("ANK-GGS-3", "GGS - 3, Ankleshwar", "Surface"),
    "GGS - 4": ("ANK-GGS-4", "GGS - 4, Ankleshwar", "Surface"),
    "GGS - 5": ("ANK-GGS-5", "GGS - 5, Ankleshwar", "Surface"),
    "GGS - 6": ("ANK-GGS-6", "GGS - 6, Ankleshwar", "Surface"),
    "GGS 253": ("ANK-GGS-253", "GGS 253, Ankleshwar", "Surface"),
    "CTF - Ank.": ("CTF-ANK", "CTF Ankleshwar", "Surface"),
    "MPH ANK": ("MPH-ANK", "MPH Ankleshwar", "Surface"),
    "GCS Motwan": ("GCS-MOTWAN", "GCS Motwan", "Surface"),
    "GGS Motwan": ("GGS-MOTWAN", "GGS Motwan", "Surface"),
    "GCS Olpad": ("GCS-OLPAD", "GCS Olpad", "Surface"),
    "KIM EPS": ("EPS-KIM", "KIM EPS", "Surface"),
    "NADA": ("GGS-NADA", "NADA GGS", "Surface"),
    "Intake Well Kathor": ("INTAKE-KATHOR", "Intake Well Kathor", "Surface"),
    "ZANORE WTP": ("WTP-ZANORE", "Zanore WTP", "Surface"),

    # Area-3
    "GGS.1": ("GDR-GGS-1", "GGS-1, Gandhar", "Surface"),
    "GGS.2": ("GDR-GGS-2", "GGS-2, Gandhar", "Surface"),
    "GGS.3": ("GDR-GGS-3", "GGS-3, Gandhar", "Surface"),
    "GGS.4": ("GDR-GGS-4", "GGS-4, Gandhar", "Surface"),
    "GGS.5": ("GDR-GGS-5", "GGS-5, Gandhar", "Surface"),
    "GGS.6": ("GDR-GGS-6", "GGS-6, Gandhar", "Surface"),
    "GGS.7": ("GDR-GGS-7", "GGS-7, Gandhar", "Surface"),
    "GGS.8": ("GDR-GGS-8", "GGS-8, Gandhar", "Surface"),
    "DSA MULLER": ("DSA-MULLER", "DSA Muller", "Surface"),
    "GGS JOLWA": ("GGS-JOLWA", "GGS Jolwa", "Surface"),
    "GGS KOSAMBA": ("GGS-KOSAMBA", "GGS Kosamba", "Surface"),
    "GGS JAMBUSAR": ("GGS-JAMBUSAR", "GGS Jambusar", "Surface"),
    "GGS.- DAHEJ": ("GGS-DAHEJ", "GGS Dahej", "Surface"),

    # Area-4
    "GGS DABKA": ("GGS-DABKA", "GGS Dabka", "Surface"),
    "GGS GNAQ GANDHAR": ("GGS-GNAQ", "GGS GNAQ Gandhar", "Surface"),

    # CPF Gandhar
    "CPF GANDHAR": ("CPF-GANDHAR", "CPF Gandhar", "Surface"),
    "CCPP, CPF GANDHAR": ("CCPP-GANDHAR", "CCPP, CPF Gandhar", "Surface"),

    # Workover Rig
    "CW-50-VII": ("RIG-CW50-VII", "CW-50-VII", "Workover Rig"),
    "AHWR-50-03": ("RIG-AHWR50-03", "AHWR-50-03", "Workover Rig"),
    "AHWR-50-06": ("RIG-AHWR50-06", "AHWR-50-06", "Workover Rig"),
    "RG-150-I": ("RIG-RG150-I", "RG-150-I", "Workover Rig"),
    "ROM-100-I": ("RIG-ROM100-I", "ROM-100-I", "Workover Rig"),
    "ROM-100-II": ("RIG-ROM100-II", "ROM-100-II", "Workover Rig"),

    # Instruments-only installations (SCADA / CTM)
    "ANK GGS01": ("ANK-GGS-1", "GGS-1, Ankleshwar", "Surface"),
    "ANKGGS-2": ("ANK-GGS-2", "GGS - 2, Ankleshwar", "Surface"),
    "ANKGGS03": ("ANK-GGS-3", "GGS - 3, Ankleshwar", "Surface"),
    "ANKGGS04": ("ANK-GGS-4", "GGS - 4, Ankleshwar", "Surface"),
    "ANKGGS05": ("ANK-GGS-5", "GGS - 5, Ankleshwar", "Surface"),
    "ANKGGS06": ("ANK-GGS-6", "GGS - 6, Ankleshwar", "Surface"),
    "ANK CTF": ("CTF-ANK", "CTF Ankleshwar", "Surface"),
    "ANK CTF LPG": ("CTF-ANK-LPG", "CTF Ankleshwar LPG", "Surface"),
    "ANK MOTWAN GCS": ("GCS-MOTWAN", "GCS Motwan", "Surface"),
    "ANK MOTWAN GGS": ("GGS-MOTWAN", "GGS Motwan", "Surface"),
    "CTF ANKLESHWAR": ("CTF-ANK", "CTF Ankleshwar", "Surface"),
    "MOTWAN GGS": ("GGS-MOTWAN", "GGS Motwan", "Surface"),
    "GDRGGS01": ("GDR-GGS-1", "GGS-1, Gandhar", "Surface"),
    "GDRGGS03": ("GDR-GGS-3", "GGS-3, Gandhar", "Surface"),
    "GDRGGS04": ("GDR-GGS-4", "GGS-4, Gandhar", "Surface"),
    "GDRGGS05": ("GDR-GGS-5", "GGS-5, Gandhar", "Surface"),
    "GDR GGS07": ("GDR-GGS-7", "GGS-7, Gandhar", "Surface"),
    "GDR 06": ("GDR-GGS-6", "GGS-6, Gandhar", "Surface"),
    "DABKA GGS": ("GGS-DABKA", "GGS Dabka", "Surface"),
    "DAHEJ GGS": ("GGS-DAHEJ", "GGS Dahej", "Surface"),
    "GNAQ GGS": ("GGS-GNAQ", "GGS GNAQ Gandhar", "Surface"),
    "GNAQ": ("GGS-GNAQ", "GGS GNAQ Gandhar", "Surface"),
    "JOLWA GGS": ("GGS-JOLWA", "GGS Jolwa", "Surface"),
    "KOSAMBA GGS": ("GGS-KOSAMBA", "GGS Kosamba", "Surface"),
    "JAMBUSAR GGS": ("GGS-JAMBUSAR", "GGS Jambusar", "Surface"),
    "NADA GGS": ("GGS-NADA", "NADA GGS", "Surface"),
    "OLPAD GCS": ("GCS-OLPAD", "GCS Olpad", "Surface"),
    "OLPAD GGS": ("GCS-OLPAD", "GCS Olpad", "Surface"),
    "EPS-253": ("ANK-GGS-253", "GGS 253, Ankleshwar", "Surface"),
    "EPS -253": ("ANK-GGS-253", "GGS 253, Ankleshwar", "Surface"),
    "GGS -4 GANDHAR": ("GDR-GGS-4", "GGS-4, Gandhar", "Surface"),
    "GRMS": ("GRMS", "GRMS", "Surface"),
    "OPDC": ("OPDC", "OPDC", "Surface"),
}


def normalize_installation(raw_name):
    """Return (installationId, location, type) tuple."""
    if not raw_name:
        return ("UNKNOWN", "Unknown", "Surface")
    key = raw_name.strip()
    if key in INSTALLATION_MAP:
        return INSTALLATION_MAP[key]
    # Fallback: slugify
    slug = re.sub(r'[^A-Za-z0-9]+', '-', key).strip('-').upper()
    return (slug, key, "Surface")


def make_equip_tag(install_id, equip_name, seen_tags):
    """Generate a unique equipment tag."""
    # Shorten equipment name
    name = equip_name.upper()
    name = re.sub(r'[^A-Z0-9 ]+', '', name)
    words = name.split()
    # Take first 3-4 significant words
    short_parts = []
    for w in words[:4]:
        if len(w) <= 3:
            short_parts.append(w)
        else:
            short_parts.append(w[:4])
    short = '-'.join(short_parts) if short_parts else 'EQUIP'
    base_tag = f"{install_id}-{short}"
    tag = base_tag
    counter = 1
    while tag in seen_tags:
        counter += 1
        tag = f"{base_tag}-{counter}"
    seen_tags.add(tag)
    return tag


def escape_ts(s):
    """Escape a string for TypeScript single-quoted string literal."""
    if s is None:
        return ""
    s = str(s)
    s = s.replace("\\", "\\\\").replace("'", "\\'").replace("\n", "\\n").replace("\r", "")
    return s


# ---------------------------------------------------------------------------
# Parse Equipment Status
# ---------------------------------------------------------------------------
def parse_equipment_status():
    wb = openpyxl.load_workbook(EQUIPMENT_FILE, data_only=True)
    equipment_rows = []

    for sname in wb.sheetnames:
        ws = wb[sname]

        if sname == "Workover Rig":
            current_install = None
            for row_idx in range(5, ws.max_row + 1):
                vals = [ws.cell(row=row_idx, column=c).value for c in range(1, 10)]
                sno, install, equip, sap_id, make, model, serial, status, reason = vals
                if install:
                    current_install = str(install).strip()
                if not equip:
                    continue
                equip_str = str(equip).strip()
                if not equip_str:
                    continue
                equipment_rows.append({
                    'sheet': sname,
                    'installation_raw': current_install,
                    'equipment': equip_str,
                    'make': str(make).strip() if make else '',
                    'model': str(model).strip() if model else '',
                    'serial': str(serial).strip() if serial else '',
                    'status': str(status).strip() if status else '',
                    'reason': str(reason).strip() if reason else '',
                    'sap_id': str(sap_id).strip() if sap_id else '',
                })
        else:
            current_install = None
            for row_idx in range(3, ws.max_row + 1):
                vals = [ws.cell(row=row_idx, column=c).value for c in range(1, 8)]
                install, equip, make, model, serial, status, reason = vals
                if install:
                    current_install = str(install).strip()
                if not equip:
                    continue
                equip_str = str(equip).strip()
                if not equip_str:
                    continue
                equipment_rows.append({
                    'sheet': sname,
                    'installation_raw': current_install,
                    'equipment': equip_str,
                    'make': str(make).strip() if make else '',
                    'model': str(model).strip() if model else '',
                    'serial': str(serial).strip() if serial else '',
                    'status': str(status).strip() if status else '',
                    'reason': str(reason).strip() if reason else '',
                    'sap_id': '',
                })

    return equipment_rows


# ---------------------------------------------------------------------------
# Parse Instruments
# ---------------------------------------------------------------------------
def parse_scada_flow_meters():
    wb = openpyxl.load_workbook(INSTRUMENTS_FILE, data_only=True)
    ws = wb["SCADA(GAS) Flow meters"]
    rows = []
    current_install = None

    for row_idx in range(1, ws.max_row + 1):
        vals = [ws.cell(row=row_idx, column=c).value for c in range(1, 11)]
        sr, install, tag, service, working, pipe_id, orifice, range_val, pt_range, scada = vals

        if install:
            install_str = str(install).strip()
            if install_str.startswith("AREA") or install_str in ("Installation Site", "Installation Site "):
                continue
            if install_str == "Location":
                continue
            current_install = install_str

        if not tag:
            continue
        tag_str = str(tag).strip()
        if not tag_str or tag_str in ("Tag No.", " Tag No."):
            continue

        rows.append({
            'installation_raw': current_install,
            'tag': tag_str,
            'service': str(service).strip() if service else '',
            'working': str(working).strip() if working else '',
            'pipe_id': str(pipe_id).strip() if pipe_id else '',
            'orifice': str(orifice).strip() if orifice else '',
            'range': str(range_val).strip() if range_val else '',
            'pt_range': str(pt_range).strip() if pt_range else '',
            'scada': str(scada).strip() if scada else '',
        })

    return rows


def parse_ctm():
    wb = openpyxl.load_workbook(INSTRUMENTS_FILE, data_only=True)
    ws = wb["CTM"]
    rows = []

    for row_idx in range(2, ws.max_row + 1):
        vals = [ws.cell(row=row_idx, column=c).value for c in range(1, 10)]
        sl, consumer, install, make, working, calibrated, scada, sap_transfer, remarks = vals

        # Skip headers and area markers
        if consumer:
            consumer_str = str(consumer).strip()
            if consumer_str.startswith("AREA") or consumer_str.startswith("SL") or consumer_str == "NAME OF CONSUMER":
                continue
        else:
            continue

        if not consumer_str:
            continue

        install_str = str(install).strip() if install else ''
        rows.append({
            'consumer': consumer_str,
            'installation_raw': install_str,
            'make': str(make).strip() if make else '',
            'working': str(working).strip() if working else '',
            'calibrated': str(calibrated).strip() if calibrated else '',
            'scada': str(scada).strip() if scada else '',
            'sap_transfer': str(sap_transfer).strip() if sap_transfer else '',
            'remarks': str(remarks).strip() if remarks else '',
        })

    return rows


# ---------------------------------------------------------------------------
# Generate TypeScript
# ---------------------------------------------------------------------------
def generate_seed_ts():
    print("Parsing Equipment Status...")
    equipment_rows = parse_equipment_status()
    print(f"  -> {len(equipment_rows)} equipment rows")

    print("Parsing SCADA Flow Meters...")
    scada_rows = parse_scada_flow_meters()
    print(f"  -> {len(scada_rows)} SCADA rows")

    print("Parsing CTM...")
    ctm_rows = parse_ctm()
    print(f"  -> {len(ctm_rows)} CTM rows")

    # -----------------------------------------------------------------------
    # Collect unique installations
    # -----------------------------------------------------------------------
    installations = OrderedDict()  # installationId -> {id, installationId, location, type}

    def ensure_installation(raw_name):
        inst_id, location, inst_type = normalize_installation(raw_name)
        if inst_id not in installations:
            installations[inst_id] = {
                'id': str(uuid.uuid4()),
                'installationId': inst_id,
                'location': location,
                'type': inst_type,
                'isActive': True,
            }
        return inst_id

    # From equipment
    for row in equipment_rows:
        ensure_installation(row['installation_raw'])
    # From SCADA
    for row in scada_rows:
        ensure_installation(row['installation_raw'])
    # From CTM
    for row in ctm_rows:
        ensure_installation(row['installation_raw'])

    print(f"\nTotal unique installations: {len(installations)}")

    # -----------------------------------------------------------------------
    # Build equipment records
    # -----------------------------------------------------------------------
    seen_tags = set()
    equipment_records = []
    for row in equipment_rows:
        inst_id = normalize_installation(row['installation_raw'])[0]
        tag = make_equip_tag(inst_id, row['equipment'], seen_tags)
        status_lower = row['status'].lower().strip()
        running_status = 'Working' if status_lower in ('working', 'running') else row['status'] if row['status'] else 'Unknown'

        equipment_records.append({
            'equipmentTag': tag,
            'category': 'RUNNING',
            'description': row['equipment'],
            'make': row['make'],
            'model': row['model'],
            'installationId': inst_id,
            'specifications': {
                'serialNo': row['serial'],
                'runningStatus': running_status,
                'reason': row['reason'],
                'sapId': row['sap_id'],
            },
        })

    print(f"Total equipment records: {len(equipment_records)}")

    # -----------------------------------------------------------------------
    # Build instrument records (SCADA flow meters)
    # -----------------------------------------------------------------------
    instrument_records = []
    seen_instrument_tags = set()

    for row in scada_rows:
        inst_id = normalize_installation(row['installation_raw'])[0]
        tag = row['tag']
        # Ensure unique tag
        if tag in seen_instrument_tags:
            base = tag
            cnt = 2
            while tag in seen_instrument_tags:
                tag = f"{base}-{cnt}"
                cnt += 1
        seen_instrument_tags.add(tag)

        desc_parts = []
        if row['service']:
            desc_parts.append(row['service'])
        if row['pipe_id'] and row['pipe_id'] != 'None':
            desc_parts.append(f"Pipe: {row['pipe_id']}")
        description = ', '.join(desc_parts) if desc_parts else f"Flow Meter {tag}"

        instrument_records.append({
            'tagId': tag,
            'type': 'FLOW_METER',
            'serviceLine': row['service'],
            'description': description,
            'make': '',
            'installationId': inst_id,
            'calibrationFreqMonths': 12,
            'specifications': {
                'pipeId': row['pipe_id'],
                'orifice': row['orifice'],
                'range': row['range'],
                'ptRange': row['pt_range'],
                'scadaConnectivity': row['scada'],
                'workingStatus': row['working'],
            },
        })

    # CTM instruments
    for row in ctm_rows:
        inst_id = normalize_installation(row['installation_raw'])[0]
        tag = f"CTM-{inst_id}-{row['consumer'][:20].upper()}"
        tag = re.sub(r'[^A-Z0-9-]', '', tag)
        if tag in seen_instrument_tags:
            base = tag
            cnt = 2
            while tag in seen_instrument_tags:
                tag = f"{base}-{cnt}"
                cnt += 1
        seen_instrument_tags.add(tag)

        instrument_records.append({
            'tagId': tag,
            'type': 'CTM',
            'serviceLine': '',
            'description': f"CTM - {row['consumer']}",
            'make': row['make'],
            'installationId': inst_id,
            'calibrationFreqMonths': 12,
            'specifications': {
                'workingStatus': row['working'],
                'calibrated': row['calibrated'],
                'scadaConnectivity': row['scada'],
                'sapDataTransfer': row['sap_transfer'],
                'remarks': row['remarks'],
            },
        })

    print(f"Total instrument records: {len(instrument_records)}")

    # -----------------------------------------------------------------------
    # Build CTM records
    # -----------------------------------------------------------------------
    ctm_records = []
    for row in ctm_rows:
        meter_id = f"CTM-{normalize_installation(row['installation_raw'])[0]}-{row['consumer'][:20].upper()}"
        meter_id = re.sub(r'[^A-Z0-9-]', '', meter_id)

        ctm_records.append({
            'id': str(uuid.uuid4()),
            'meterId': meter_id,
            'customerName': row['consumer'],
            'meterType': row['make'],
            'product': 'GAS',
            'fluidType': 'GAS',
        })

    print(f"Total CTM records: {len(ctm_records)}")

    # -----------------------------------------------------------------------
    # Write TypeScript seed file
    # -----------------------------------------------------------------------
    lines = []
    lines.append("import { PrismaClient } from '@prisma/client';")
    lines.append("")
    lines.append("const prisma = new PrismaClient();")
    lines.append("")

    # Installations array
    lines.append("const installations = [")
    for inst in installations.values():
        lines.append("  {")
        lines.append(f"    id: '{escape_ts(inst['id'])}',")
        lines.append(f"    installationId: '{escape_ts(inst['installationId'])}',")
        lines.append(f"    location: '{escape_ts(inst['location'])}',")
        lines.append(f"    type: '{escape_ts(inst['type'])}',")
        lines.append(f"    isActive: true,")
        lines.append("  },")
    lines.append("];")
    lines.append("")

    # Equipment array
    lines.append("const equipmentList = [")
    for rec in equipment_records:
        specs_json = json.dumps(rec['specifications'])
        lines.append("  {")
        lines.append(f"    equipmentTag: '{escape_ts(rec['equipmentTag'])}',")
        lines.append(f"    category: 'RUNNING',")
        lines.append(f"    description: '{escape_ts(rec['description'])}',")
        lines.append(f"    make: '{escape_ts(rec['make'])}',")
        lines.append(f"    model: '{escape_ts(rec['model'])}',")
        lines.append(f"    installationId: '{escape_ts(rec['installationId'])}',")
        lines.append(f"    specifications: {specs_json},")
        lines.append("  },")
    lines.append("];")
    lines.append("")

    # Instruments array
    lines.append("const instrumentList = [")
    for rec in instrument_records:
        specs_json = json.dumps(rec['specifications'])
        lines.append("  {")
        lines.append(f"    tagId: '{escape_ts(rec['tagId'])}',")
        lines.append(f"    type: '{escape_ts(rec['type'])}',")
        lines.append(f"    serviceLine: '{escape_ts(rec['serviceLine'])}',")
        lines.append(f"    description: '{escape_ts(rec['description'])}',")
        lines.append(f"    make: '{escape_ts(rec['make'])}',")
        lines.append(f"    installationId: '{escape_ts(rec['installationId'])}',")
        lines.append(f"    calibrationFreqMonths: {rec['calibrationFreqMonths']},")
        lines.append(f"    specifications: {specs_json},")
        lines.append("  },")
    lines.append("];")
    lines.append("")

    # CTM array
    lines.append("const ctmList = [")
    for rec in ctm_records:
        lines.append("  {")
        lines.append(f"    id: '{escape_ts(rec['id'])}',")
        lines.append(f"    meterId: '{escape_ts(rec['meterId'])}',")
        lines.append(f"    customerName: '{escape_ts(rec['customerName'])}',")
        lines.append(f"    meterType: '{escape_ts(rec['meterType'])}',")
        lines.append(f"    product: 'GAS',")
        lines.append(f"    fluidType: 'GAS',")
        lines.append("  },")
    lines.append("];")
    lines.append("")

    # Main seeding function
    lines.append("async function main() {")
    lines.append("  console.log('Starting database seed from Excel data...');")
    lines.append("")
    lines.append("  // --- Seed Installations ---")
    lines.append("  console.log(`Seeding ${installations.length} installations...`);")
    lines.append("  for (const inst of installations) {")
    lines.append("    await prisma.installation.upsert({")
    lines.append("      where: { installationId: inst.installationId },")
    lines.append("      update: {},")
    lines.append("      create: {")
    lines.append("        id: inst.id,")
    lines.append("        installationId: inst.installationId,")
    lines.append("        location: inst.location,")
    lines.append("        type: inst.type,")
    lines.append("        isActive: inst.isActive,")
    lines.append("      },")
    lines.append("    });")
    lines.append("  }")
    lines.append("  console.log('  Installations seeded.');")
    lines.append("")

    lines.append("  // --- Build installation lookup (installationId -> DB id) ---")
    lines.append("  const allInstallations = await prisma.installation.findMany();")
    lines.append("  const installLookup: Record<string, string> = {};")
    lines.append("  for (const i of allInstallations) {")
    lines.append("    installLookup[i.installationId] = i.id;")
    lines.append("  }")
    lines.append("")

    lines.append("  // --- Seed Equipment ---")
    lines.append("  console.log(`Seeding ${equipmentList.length} equipment records...`);")
    lines.append("  for (const eq of equipmentList) {")
    lines.append("    const dbInstallId = installLookup[eq.installationId];")
    lines.append("    if (!dbInstallId) {")
    lines.append("      console.warn(`  Skipping equipment '${eq.equipmentTag}': installation '${eq.installationId}' not found`);")
    lines.append("      continue;")
    lines.append("    }")
    lines.append("    await prisma.runningEquipmentMaster.upsert({")
    lines.append("      where: { equipmentTag: eq.equipmentTag },")
    lines.append("      update: {},")
    lines.append("      create: {")
    lines.append("        equipmentTag: eq.equipmentTag,")
    lines.append("        category: eq.category,")
    lines.append("        description: eq.description,")
    lines.append("        make: eq.make,")
    lines.append("        model: eq.model,")
    lines.append("        installationId: dbInstallId,")
    lines.append("        specifications: eq.specifications as any,")
    lines.append("      },")
    lines.append("    });")
    lines.append("  }")
    lines.append("  console.log('  Equipment seeded.');")
    lines.append("")

    lines.append("  // --- Seed Instruments ---")
    lines.append("  console.log(`Seeding ${instrumentList.length} instrument records...`);")
    lines.append("  for (const inst of instrumentList) {")
    lines.append("    const dbInstallId = installLookup[inst.installationId];")
    lines.append("    if (!dbInstallId) {")
    lines.append("      console.warn(`  Skipping instrument '${inst.tagId}': installation '${inst.installationId}' not found`);")
    lines.append("      continue;")
    lines.append("    }")
    lines.append("    await prisma.instrumentMaster.upsert({")
    lines.append("      where: { tagId: inst.tagId },")
    lines.append("      update: {},")
    lines.append("      create: {")
    lines.append("        tagId: inst.tagId,")
    lines.append("        type: inst.type,")
    lines.append("        serviceLine: inst.serviceLine,")
    lines.append("        description: inst.description,")
    lines.append("        make: inst.make,")
    lines.append("        installationId: dbInstallId,")
    lines.append("        calibrationFreqMonths: inst.calibrationFreqMonths,")
    lines.append("      },")
    lines.append("    });")
    lines.append("  }")
    lines.append("  console.log('  Instruments seeded.');")
    lines.append("")

    lines.append("  // --- Seed Custody Transfer Meters ---")
    lines.append("  console.log(`Seeding ${ctmList.length} CTM records...`);")
    lines.append("  for (const ctm of ctmList) {")
    lines.append("    await prisma.custodyTransferMeter.upsert({")
    lines.append("      where: { meterId: ctm.meterId },")
    lines.append("      update: {},")
    lines.append("      create: {")
    lines.append("        id: ctm.id,")
    lines.append("        meterId: ctm.meterId,")
    lines.append("        customerName: ctm.customerName,")
    lines.append("        meterType: ctm.meterType,")
    lines.append("        product: ctm.product,")
    lines.append("        fluidType: ctm.fluidType,")
    lines.append("      },")
    lines.append("    });")
    lines.append("  }")
    lines.append("  console.log('  CTM records seeded.');")
    lines.append("")

    lines.append("  console.log('Database seeding complete!');")
    lines.append("}")
    lines.append("")
    lines.append("main()")
    lines.append("  .then(() => prisma.$disconnect())")
    lines.append("  .catch((e) => {")
    lines.append("    console.error('Seed error:', e);")
    lines.append("    prisma.$disconnect();")
    lines.append("    process.exit(1);")
    lines.append("  });")

    ts_content = "\n".join(lines) + "\n"

    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        f.write(ts_content)

    print(f"\nSeed file written to: {OUTPUT_FILE}")
    print(f"  {len(installations)} installations")
    print(f"  {len(equipment_records)} equipment")
    print(f"  {len(instrument_records)} instruments")
    print(f"  {len(ctm_records)} CTM records")


if __name__ == '__main__':
    generate_seed_ts()
